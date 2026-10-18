"""Import the 12 October 2026 BLUB, Elite and Funkier stock workbooks.

Workbook stock is authoritative. Blank, zero and error stock values are listed as
Coming Soon. Dealer prices and volume offers are deliberately ignored.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "products" / "distributor"
DATA = ROOT / "assets" / "data" / "distributor-products.json"
ONLINE_IMAGES = json.loads((ROOT / "assets" / "data" / "online-product-images.json").read_text(encoding="utf-8"))
SOURCES = {
    "blub": Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"E:\Downloads\BLUB LUBES stock sheet   12-10-2026.xlsx"),
    "elite": Path(sys.argv[2]) if len(sys.argv) > 2 else Path(r"E:\Downloads\Elite stock sheet 12-10-2026.xlsx"),
    "funkier": Path(sys.argv[3]) if len(sys.argv) > 3 else Path(r"E:\Downloads\Funkier Never Before Price  12-10-2026.xlsx"),
}
ELITE_PRICE_LOOKUP = {
    "0165001": (99990, 99990, "R R Cycles and Manoj Cycles India, checked 2026-10-18"),
    "0172001": (29900, 34900, "Existing Cyclify/CycleTime listing, checked 2026-10-18"),
    "0192003P": (33990, 46990, "Manoj Cycles India, checked 2026-10-18"),
    "121026": (30370, 30990, "Decathlon India, checked 2026-10-18"),
    "0121006": (29990, 29990, "Cobbled Climbs India, checked 2026-10-18"),
}


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def stock(value):
    try:
        amount = int(float(value))
        return amount if amount > 0 else 0
    except (TypeError, ValueError):
        return 0


def gtin(value):
    text = clean(value).replace(".0", "")
    return text if text.isdigit() and len(text) in (8, 12, 13, 14) else None


def key_for(*parts):
    return hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest()[:16]


def embedded_images(sheet):
    result = defaultdict(list)
    for image in sheet._images:
        row, col = image.anchor._from.row + 1, image.anchor._from.col + 1
        result[(row, col)].append(image._data())
    return result


def save_image(payload, label):
    digest = hashlib.sha256(payload).hexdigest()[:20]
    rel = Path("assets/products/distributor") / f"{digest}.webp"
    thumb_rel = Path("assets/products/distributor/thumbs") / f"{digest}.webp"
    full, thumb = ROOT / rel, ROOT / thumb_rel
    full.parent.mkdir(parents=True, exist_ok=True)
    thumb.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(io.BytesIO(payload)) as source:
        source = source.convert("RGB")
        source.thumbnail((1500, 1700), Image.Resampling.LANCZOS)
        source.save(full, "WEBP", quality=90, method=6)
        preview = source.copy()
        preview.thumbnail((360, 400), Image.Resampling.LANCZOS)
        preview.save(thumb, "WEBP", quality=82, method=6)
    return rel.as_posix()


def category_for_elite(sheet_category, name):
    text = f"{sheet_category} {name}".lower()
    if "bottle" in text or "cage" in text or "drinking" in text or "skekane" in text:
        return "accessories.html", "Accessories", "bottle-cages"
    return "electronics.html", "Electronics", "smart-trainer"


def product(*, product_id, name, brand, model, price, mrp, quantity, image_paths,
            category_page, category_label, main_category, description, specs, ean=None,
            sizes=None, colours=None, variants=None, price_label=None):
    available = quantity > 0
    item = {
        "id": product_id,
        "name": name,
        "brand": brand,
        "inventoryGroup": "expanded",
        "inventorySource": "distributor-stock-20261012",
        "model": model,
        "mpn": model,
        "priceLabel": price_label or "MRP",
        "primaryLabel": "Select Size" if sizes else "Quantity",
        "categoryPage": category_page,
        "categoryLabel": category_label,
        "mainCategory": main_category,
        "subCategory": brand.lower(),
        "price": int(price),
        "mrp": int(mrp),
        "available": available,
        "stockStatus": "available" if available else "coming-soon",
        "stockQuantity": quantity,
        "image": image_paths[0],
        "images": list(dict.fromkeys(image_paths)),
        "description": description,
        "keywords": f"{brand} {model} {name} India buy online cycling",
        "specs": specs,
    }
    if ean:
        item["gtin"] = ean
    if sizes:
        item.update({"sizes": sizes, "secondaryLabel": "Select Colour", "secondaryOptions": colours or [],
                     "purchaseVariants": variants or [], "defaultVariantKey": next((v["key"] for v in variants or [] if v["available"]), (variants or [{}])[0].get("key"))})
    return item


def import_blub(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=True).active
    images = embedded_images(sheet)
    for row in range(3, 27):
        name, model = clean(sheet.cell(row, 5).value), clean(sheet.cell(row, 4).value)
        if not name:
            continue
        image_paths = [save_image(raw, f"blub-{model}") for raw in images[(row, 8)]]
        if not image_paths:
            audit["held"].append({"brand": "BLUB", "model": model, "reason": "No workbook product image"})
            continue
        quantity = stock(sheet.cell(row, 6).value)
        mrp = int(sheet.cell(row, 7).value)
        category = clean(sheet.cell(row, 3).value)
        products.append(product(product_id=920000 + row, name=name, brand="BLUB", model=model,
            price=mrp, mrp=mrp, quantity=quantity, image_paths=image_paths,
            category_page="components.html", category_label="Components", main_category="maintenance",
            description=f"{name} for bicycle care and maintenance. Stock status is based on the BLUB distributor sheet dated 12 October 2026.",
            specs=["Brand : BLUB", f"Model : {model}", f"Category : {category}"], ean=gtin(sheet.cell(row, 2).value)))


def import_elite(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=True).active
    images = embedded_images(sheet)
    for row in range(3, 67):
        raw_name = clean(sheet.cell(row, 2).value)
        if not raw_name:
            continue
        model = clean(sheet.cell(row, 5).value)
        image_paths = [save_image(raw, f"elite-{model}") for raw in images[(row, 10)]]
        sourced = ONLINE_IMAGES.get(f"Elite|{model}")
        if not image_paths and sourced:
            image_paths = [sourced["image"]]
            audit.setdefault("sourcedImages", []).append({"brand": "Elite", "model": model, "sourcePage": sourced["sourcePage"], "sourceName": sourced["sourceName"]})
        if not image_paths:
            audit["held"].append({"brand": "Elite", "model": model, "name": raw_name, "reason": "No workbook product image"})
            continue
        supplied_mrp = sheet.cell(row, 8).value
        if supplied_mrp is None:
            price, mrp, source = ELITE_PRICE_LOOKUP[model]
            audit["researchedPrices"].append({"model": model, "price": price, "mrp": mrp, "source": source})
        else:
            price = mrp = int(supplied_mrp)
        quantity = stock(sheet.cell(row, 9).value)
        name = "Elite " + re.sub(r"\s+", " ", raw_name).strip().title()
        name = name.replace("Drivo With Bag Vaisa Free!", "Drivo Trainer with Vaisa Bag")
        if model == "0172001":
            name = "Elite Turno Home Trainer"
        page, label, main = category_for_elite(sheet.cell(row, 6).value, raw_name)
        details = clean(sheet.cell(row, 7).value)
        specs = ["Brand : Elite", f"Model : {model}", f"Category : {clean(sheet.cell(row, 6).value)}"]
        specs.extend([clean(part) for part in re.split(r"[\n\r]+", details) if clean(part)][:8])
        product_id = 10994532548801 if model == "0172001" else 921000 + row
        products.append(product(product_id=product_id, name=name, brand="Elite", model=model,
            price=price, mrp=mrp, quantity=quantity, image_paths=image_paths,
            category_page=page, category_label=label, main_category=main,
            description=f"{name} from Elite. Stock status is based on the Elite distributor sheet dated 12 October 2026.",
            specs=specs, ean=gtin(sheet.cell(row, 4).value), price_label="SPECIAL PRICE" if price < mrp else "MRP"))


def import_funkier(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=True).active
    images = embedded_images(sheet)
    starts = [row for row in range(3, 49) if sheet.cell(row, 2).value]
    grouped = defaultdict(lambda: {"variants": [], "images": [], "spec": "", "composition": ""})
    for index, start in enumerate(starts):
        end = (starts[index + 1] - 1) if index + 1 < len(starts) else sheet.max_row
        model = clean(sheet.cell(start, 2).value)
        group = grouped[model]
        group["spec"] = group["spec"] or clean(sheet.cell(start, 4).value)
        group["composition"] = group["composition"] or clean(sheet.cell(start, 5).value)
        group["images"].extend(save_image(raw, f"funkier-{model}") for col in (12, 13) for raw in images[(start, col)])
        for row in range(start, end + 1):
            if not sheet.cell(row, 3).value or not sheet.cell(row, 6).value or not sheet.cell(row, 7).value:
                continue
            group["variants"].append({
                "ean": gtin(sheet.cell(row, 3).value), "colour": clean(sheet.cell(row, 6).value),
                "size": clean(sheet.cell(row, 7).value), "stock": stock(sheet.cell(row, 8).value),
                "mrp": int(sheet.cell(row, 9).value), "price": int(sheet.cell(row, 10).value),
            })
    for offset, (model, group) in enumerate(grouped.items(), 1):
        if not group["images"] or not group["variants"]:
            audit["held"].append({"brand": "Funkier", "model": model, "reason": "Missing image or sale variant"})
            continue
        variants = []
        for variant in group["variants"]:
            variant_key = key_for("funkier", model, variant["colour"], variant["size"])
            variants.append({"key": variant_key, "size": variant["size"], "colour": variant["colour"], "build": "",
                "price": variant["price"], "mrp": variant["mrp"], "available": variant["stock"] > 0,
                "stockStatus": "available" if variant["stock"] > 0 else "coming-soon", "image": group["images"][0],
                **({"gtin": variant["ean"]} if variant["ean"] else {})})
        in_stock = [v for v in group["variants"] if v["stock"] > 0]
        chosen = min(in_stock or group["variants"], key=lambda value: value["price"])
        name = f"Funkier {group['spec'] or model}"
        products.append(product(product_id=922000 + offset, name=name, brand="Funkier", model=model,
            price=chosen["price"], mrp=chosen["mrp"], quantity=sum(v["stock"] for v in group["variants"]),
            image_paths=group["images"], category_page="wearables.html", category_label="Apparel & Wearables",
            main_category="jersey-shorts", description=f"{name}. Choose the available size and colour before checkout.",
            specs=["Brand : Funkier", f"Model : {model}", f"Material : {group['composition']}", f"Product : {group['spec']}"],
            sizes=list(dict.fromkeys(v["size"] for v in group["variants"])), colours=list(dict.fromkeys(v["colour"] for v in group["variants"])),
            variants=variants, price_label="NEVER BEFORE PRICE"))


def main():
    missing = [str(path) for path in SOURCES.values() if not path.exists()]
    if missing:
        raise SystemExit("Missing workbook(s): " + ", ".join(missing))
    products, audit = [], {"held": [], "researchedPrices": []}
    import_blub(SOURCES["blub"], products, audit)
    import_elite(SOURCES["elite"], products, audit)
    import_funkier(SOURCES["funkier"], products, audit)
    payload = {"version": 1, "stockAsOf": "2026-10-12", "products": products, "audit": audit}
    DATA.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    allowlist_path = ROOT / "assets" / "data" / "product-allowlist.json"
    allowlist = json.loads(allowlist_path.read_text(encoding="utf-8"))
    allowlist["productIds"] = sorted(set(map(str, allowlist["productIds"])) | {str(p["id"]) for p in products}, key=int)
    allowlist_path.write_text(json.dumps(allowlist, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"products": len(products), "available": sum(p["available"] for p in products),
                      "comingSoon": sum(not p["available"] for p in products), "held": audit["held"],
                      "researchedPrices": audit["researchedPrices"]}, indent=2))


if __name__ == "__main__":
    main()
