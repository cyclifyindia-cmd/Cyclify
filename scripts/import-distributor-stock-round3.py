"""Import six October 2026 distributor workbooks into the Cyclify catalogue."""
from __future__ import annotations

import hashlib
import html
import io
import json
import re
import sys
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path

import openpyxl
import xlrd
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "products" / "distributor-3"
DATA = ROOT / "assets" / "data" / "distributor-products-3.json"
CACHE_PATH = ROOT / "assets" / "data" / "distributor-3-image-cache.json"
SOURCE_CACHE = json.loads(CACHE_PATH.read_text(encoding="utf-8")) if CACHE_PATH.exists() else {}
UA = {"User-Agent": "Mozilla/5.0 (compatible; CyclifyCatalog/1.0; +https://cyclify.in/)"}
SOURCES = {
    "tifosi": Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"E:\Downloads\Tifosi Stock Order Sheet (new arrival) Oct.26.xlsx"),
    "ravemen": Path(sys.argv[2]) if len(sys.argv) > 2 else Path(r"E:\Downloads\Ravemen Light stock sheet 12-10-2026.xlsx"),
    "safety": Path(sys.argv[3]) if len(sys.argv) > 3 else Path(r"E:\Downloads\SAFETY LABS Helmets stock sheet   12-10-2026.xlsx"),
    "tacx": Path(sys.argv[4]) if len(sys.argv) > 4 else Path(r"E:\Downloads\TACX Trainers & Rollers stock sheet  12-10-2026.xlsx"),
    "polygon": Path(sys.argv[5]) if len(sys.argv) > 5 else Path(r"E:\Downloads\Polygon Stock List 16th Oct.xls"),
    "merida": Path(sys.argv[6]) if len(sys.argv) > 6 else Path(r"E:\Downloads\Merida current stock  12-Oct-2026.xlsx"),
}


def clean(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value or "")).strip()


def norm(value):
    return re.sub(r"[^a-z0-9]+", " ", clean(value).casefold()).strip()


def money(value):
    try:
        result = int(round(float(value)))
        return result if result > 0 else 0
    except (TypeError, ValueError):
        return 0


def quantity(value):
    try:
        result = int(float(value))
        return result if result > 0 else 0
    except (TypeError, ValueError):
        text = clean(value).casefold()
        match = re.search(r"(?:only\s*)?(\d+)", text)
        if match:
            return int(match.group(1))
        return 1 if text in {"available", "avaiable", "in stock"} else 0


def gtin(value):
    text = re.sub(r"\.0$", "", clean(value))
    return text if text.isdigit() and len(text) in (8, 12, 13, 14) else None


def key_for(*parts):
    return hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest()[:16]


def fetch(url, timeout=45):
    request = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return response.read()


def embedded_images(sheet, columns=None):
    result = defaultdict(list)
    for image in sheet._images:
        row = image.anchor._from.row + 1
        col = image.anchor._from.col + 1
        if columns is None or col in columns:
            result[row].append(image._data())
    return result


def save_image(payload):
    digest = hashlib.sha256(payload).hexdigest()[:20]
    rel = Path("assets/products/distributor-3") / f"{digest}.webp"
    thumb_rel = Path("assets/products/distributor-3/thumbs") / f"{digest}.webp"
    full, thumb = ROOT / rel, ROOT / thumb_rel
    if full.exists() and thumb.exists():
        return rel.as_posix()
    full.parent.mkdir(parents=True, exist_ok=True)
    thumb.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(io.BytesIO(payload)) as source:
        source.seek(0)
        source = source.convert("RGB")
        source.thumbnail((1500, 1700), Image.Resampling.LANCZOS)
        source.save(full, "WEBP", quality=90, method=6)
        preview = source.copy()
        preview.thumbnail((360, 400), Image.Resampling.LANCZOS)
        preview.save(thumb, "WEBP", quality=82, method=6)
    return rel.as_posix()


def save_remote(url):
    return save_image(fetch(url))


def product(*, product_id, name, brand, model, price, mrp, stock, images,
            page, label, main, sub, description, specs, variants=None,
            sizes=None, colours=None, ean=None, source_url=None):
    if not images:
        raise RuntimeError(f"No genuine photo for {brand} {name}")
    if price < 1:
        raise RuntimeError(f"No public price for {brand} {name}")
    mrp = max(mrp, price)
    available = stock > 0
    result = {
        "id": product_id, "name": name, "brand": brand,
        "inventoryGroup": "expanded", "inventorySource": "distributor-stock-202610-round3",
        "model": model, "mpn": model, "priceLabel": "OFFER PRICE" if price < mrp else "MRP",
        "primaryLabel": "Select Size" if sizes else "Quantity",
        "categoryPage": page, "categoryLabel": label, "mainCategory": main, "subCategory": sub,
        "price": price, "mrp": mrp, "available": available,
        "stockStatus": "available" if available else "coming-soon", "stockQuantity": stock,
        "image": images[0], "images": list(dict.fromkeys(images)), "description": description,
        "keywords": f"{brand} {model} {name} India buy online cycling",
        "specs": specs,
    }
    if ean:
        result["gtin"] = ean
    if source_url:
        result["manufacturerUrl"] = source_url
    if variants:
        result.update({
            "sizes": sizes or [], "secondaryLabel": "Select Colour",
            "secondaryOptions": colours or [], "purchaseVariants": variants,
            "defaultVariantKey": next((v["key"] for v in variants if v["available"]), variants[0]["key"]),
        })
    return result


def import_tifosi(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=True)["TIFOSI Order Sheet"]
    images = embedded_images(sheet, {11, 12})
    for row in range(3, 65):
        model, name = clean(sheet.cell(row, 2).value), clean(sheet.cell(row, 5).value)
        if not name:
            continue
        image_paths = [save_image(payload) for payload in images[row]]
        if not image_paths:
            raise RuntimeError(f"Tifosi workbook image missing on row {row}: {name}")
        qty, mrp = quantity(sheet.cell(row, 9).value), money(sheet.cell(row, 10).value)
        lens = clean(sheet.cell(row, 7).value)
        product_name = name if name.casefold().startswith("tifosi") else f"Tifosi {name}"
        products.append(product(
            product_id=926000 + row, name=product_name, brand="Tifosi", model=model or clean(sheet.cell(row, 3).value),
            price=mrp, mrp=mrp, stock=qty, images=image_paths, page="accessories.html",
            label="Accessories", main="eyewear", sub="cycling sunglasses",
            description=f"{product_name} cycling eyewear. Stock is from the October 2026 distributor sheet.",
            specs=["Brand : Tifosi", f"Category : {clean(sheet.cell(row, 4).value)}", f"Lens : {lens}", f"Type : {clean(sheet.cell(row, 8).value)}"],
            source_url=clean(sheet.cell(row, 15).value) or None))
    audit["workbookImages"]["Tifosi"] = len(products)


def import_ravemen(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=True)["Ravemen Stock Sheet"]
    images = embedded_images(sheet, {3})
    start = len(products)
    for row in range(3, 46):
        name = clean(sheet.cell(row, 2).value)
        if not name:
            continue
        image_paths = [save_image(payload) for payload in images[row]]
        if not image_paths:
            raise RuntimeError(f"Ravemen workbook image missing on row {row}: {name}")
        qty, mrp = quantity(sheet.cell(row, 14).value), money(sheet.cell(row, 13).value)
        model = clean(sheet.cell(row, 12).value) or name
        full_name = name if name.casefold().startswith("ravemen") else f"Ravemen {name}"
        products.append(product(
            product_id=927000 + row, name=full_name, brand="Ravemen", model=model,
            price=mrp, mrp=mrp, stock=qty, images=image_paths, page="electronics.html",
            label="Electronics", main="lights", sub=clean(sheet.cell(row, 1).value).casefold(),
            description=f"{full_name} bicycle light. Stock is from the Ravemen distributor sheet dated 12 October 2026.",
            specs=["Brand : Ravemen", f"Maximum output : {clean(sheet.cell(row, 6).value)}", f"Maximum run time : {clean(sheet.cell(row, 7).value)}", f"Weight : {clean(sheet.cell(row, 11).value)}"],
            ean=gtin(sheet.cell(row, 12).value)))
    audit["workbookImages"]["Ravemen"] = len(products) - start


def import_safety(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=True)["Safety Labs"]
    images = embedded_images(sheet, {8})
    rows = []
    last_image = None
    for row in range(3, 57):
        name = clean(sheet.cell(row, 4).value)
        if not name:
            continue
        if images[row]:
            last_image = save_image(images[row][0])
        rows.append((row, name, last_image))
    grouped = defaultdict(list)
    for row, name, image_path in rows:
        grouped[norm(name)].append((row, name, image_path))
    for offset, group in enumerate(grouped.values(), 1):
        variants, sizes, colours = [], [], []
        for row, name, image_path in group:
            size = clean(sheet.cell(row, 5).value) or "Standard"
            colour_match = re.search(r"-\s*([^-]+)$", name)
            colour = colour_match.group(1).strip() if colour_match else "Standard"
            qty, mrp = quantity(sheet.cell(row, 6).value), money(sheet.cell(row, 7).value)
            variants.append({"key": key_for("safety", name, size), "size": size, "colour": colour, "build": "", "price": mrp, "mrp": mrp, "available": qty > 0, "stockStatus": "available" if qty > 0 else "coming-soon", "stockQuantity": qty, "image": image_path, **({"gtin": gtin(sheet.cell(row, 3).value)} if gtin(sheet.cell(row, 3).value) else {})})
            sizes.append(size); colours.append(colour)
        chosen = min([v for v in variants if v["available"]] or variants, key=lambda v: v["price"])
        first_row, raw_name, image_path = group[0]
        full_name = raw_name if raw_name.casefold().startswith("safety labs") else f"Safety Labs {raw_name}"
        products.append(product(
            product_id=928000 + offset, name=full_name, brand="Safety Labs", model=clean(sheet.cell(first_row, 3).value) or raw_name,
            price=chosen["price"], mrp=chosen["mrp"], stock=sum(v["stockQuantity"] for v in variants), images=[image_path],
            page="wearables.html", label="Apparel & Wearables", main="helmets", sub="cycling helmets",
            description=f"{full_name} cycling helmet. Select an available size before checkout.",
            specs=["Brand : Safety Labs", f"Category : {clean(sheet.cell(first_row, 2).value)}"],
            variants=variants, sizes=list(dict.fromkeys(sizes)), colours=list(dict.fromkeys(colours))))
    audit["workbookImages"]["Safety Labs"] = len(grouped)


def import_tacx(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=False)["Tacx"]
    values = openpyxl.load_workbook(path, data_only=True)["Tacx"]
    images = embedded_images(sheet, {8})
    formula_prices = {6: 22990, 8: 14590}
    start = len(products)
    for row in range(3, 12):
        name = clean(values.cell(row, 2).value)
        if not name:
            continue
        image_paths = [save_image(payload) for payload in images[row]]
        if not image_paths:
            raise RuntimeError(f"Tacx workbook image missing on row {row}: {name}")
        mrp = money(values.cell(row, 6).value) or formula_prices.get(row, 0)
        qty = quantity(values.cell(row, 7).value)
        model = clean(values.cell(row, 3).value) or clean(values.cell(row, 4).value) or name
        is_trainer = any(token in name.casefold() for token in ("trainer", "roller", "boost", "flux", "neo"))
        products.append(product(
            product_id=929000 + row, name=name if name.casefold().startswith("tacx") else f"Tacx {name}", brand="Tacx", model=model,
            price=mrp, mrp=mrp, stock=qty, images=image_paths,
            page="electronics.html" if is_trainer else "accessories.html",
            label="Electronics" if is_trainer else "Accessories", main="smart-trainers" if is_trainer else "trainer-accessories",
            sub="trainers and rollers" if is_trainer else "trainer accessories",
            description=clean(values.cell(row, 5).value) or f"Tacx {name} cycling training product.",
            specs=["Brand : Tacx", f"SKU : {model}"], source_url=clean(values.cell(row, 8).value) if clean(values.cell(row, 8).value).startswith("http") else None))
    audit["workbookImages"]["Tacx"] = len(products) - start


def polygon_image(model, year, cache, audit):
    cache_key = norm(model)
    disk_key = f"Polygon|{cache_key}"
    if disk_key in SOURCE_CACHE and (ROOT / SOURCE_CACHE[disk_key]["image"]).exists():
        saved = SOURCE_CACHE[disk_key]
        audit["sourcedImages"].append({"brand": "Polygon", "model": model, "sourcePage": saved["sourcePage"], "sourceName": "Polygon official"})
        return saved["image"], saved["sourcePage"]
    if cache_key in cache:
        return cache[cache_key]
    queries = [model]
    simplified = re.sub(r"\b(?:DI2|12S|11S|10S|9S|8S|LTD|LIMITED|EDITION)\b.*$", "", model, flags=re.I).strip(" -")
    if simplified and norm(simplified) != cache_key:
        queries.append(simplified)
    candidates = []
    for query in queries:
        search = "https://www.polygonbikes.com/wp-json/wp/v2/search?" + urllib.parse.urlencode({"search": query, "per_page": 20})
        entries = json.loads(fetch(search))
        candidates = [item for item in entries if item.get("subtype") == "product"]
        if candidates:
            break
    exact_keys = {cache_key, norm(simplified)}
    exact = [item for item in candidates if norm(html.unescape(item.get("title", ""))) in exact_keys]
    candidates = exact or candidates
    if not candidates:
        slug = re.sub(r"[^a-z0-9]+", "-", simplified.casefold()).strip("-")
        family = slug.split("-")[0]
        direct_pages = [
            f"https://www.polygonbikes.com/{clean(year)}-{slug}-mountain-bike/",
            f"https://www.polygonbikes.com/{clean(year)}-{slug}-road-bike/",
            f"https://www.polygonbikes.com/{slug}/",
            f"https://www.polygonbikes.com/{slug}-2/",
            f"https://www.polygonbikes.com/us/{slug}/",
            f"https://www.polygonbikes.com/us/{clean(year)}-{slug}-mountain-bike/",
            f"https://www.polygonbikes.com/{family}/",
        ]
        for page_url in direct_pages:
            try:
                page_html = fetch(page_url, timeout=10).decode("utf-8", "ignore")
            except Exception:
                continue
            match = re.search(r'<meta[^>]+property=["\']og:image["\'][^>]+content=["\']([^"\']+)', page_html, re.I)
            if not match:
                match = re.search(r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+property=["\']og:image', page_html, re.I)
            image_url = html.unescape(match.group(1)) if match else None
            if image_url and "logo" not in image_url.casefold():
                result = (save_remote(image_url), page_url)
                cache[cache_key] = result
                SOURCE_CACHE[disk_key] = {"image": result[0], "sourcePage": result[1]}
                CACHE_PATH.write_text(json.dumps(SOURCE_CACHE, indent=2) + "\n", encoding="utf-8")
                audit["sourcedImages"].append({"brand": "Polygon", "model": model, "sourcePage": page_url, "sourceName": "Polygon official"})
                return result
        raise RuntimeError(f"No official Polygon product match for {model}")
    year_text = clean(year)
    ordered = sorted(candidates, key=lambda item: 0 if year_text and year_text in item.get("url", "") else 1)
    chosen, image_url = ordered[0], None
    for candidate in ordered:
        endpoint = f"https://www.polygonbikes.com/wp-json/wp/v2/product/{candidate['id']}?_embed=1"
        detail = json.loads(fetch(endpoint))
        image_url = detail.get("_embedded", {}).get("wp:featuredmedia", [{}])[0].get("source_url")
        if image_url:
            chosen = candidate
            break
    if not image_url:
        page = fetch(chosen["url"]).decode("utf-8", "ignore")
        match = re.search(r'<meta[^>]+property=["\']og:image["\'][^>]+content=["\']([^"\']+)', page, re.I)
        if not match:
            match = re.search(r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+property=["\']og:image', page, re.I)
        image_url = html.unescape(match.group(1)) if match else None
    if not image_url:
        raise RuntimeError(f"Official Polygon image unavailable for {model}")
    result = (save_remote(image_url), chosen["url"])
    cache[cache_key] = result
    SOURCE_CACHE[disk_key] = {"image": result[0], "sourcePage": result[1]}
    CACHE_PATH.write_text(json.dumps(SOURCE_CACHE, indent=2) + "\n", encoding="utf-8")
    audit["sourcedImages"].append({"brand": "Polygon", "model": model, "sourcePage": chosen["url"], "sourceName": "Polygon official"})
    return result


def import_polygon(path, products, audit):
    sheet = xlrd.open_workbook(path).sheet_by_name("New and Current Stock")
    grouped = defaultdict(list)
    for index in range(1, sheet.nrows):
        values = sheet.row_values(index)
        name = clean(values[2])
        if not name:
            continue
        group_key = (clean(values[0]), norm(name), norm(values[4]))
        grouped[group_key].append(values)
    image_cache = {}
    for offset, rows in enumerate(grouped.values(), 1):
        year, name, category, colour = clean(rows[0][0]), clean(rows[0][2]), clean(rows[0][3]), clean(rows[0][4]) or "Standard"
        print(f"Polygon {offset}/{len(grouped)}: {name}", flush=True)
        image_path, source_url = polygon_image(name, year, image_cache, audit)
        variants, sizes = [], []
        for row in rows:
            size = clean(row[5]) or "Standard"
            mrp = money(row[6])
            price = money(row[9]) or money(row[8]) or mrp
            qty = quantity(row[7])
            variants.append({"key": key_for("polygon", year, name, colour, size), "size": size, "colour": colour, "build": "", "price": price, "mrp": max(mrp, price), "available": qty > 0, "stockStatus": "available" if qty > 0 else "coming-soon", "stockQuantity": qty, "image": image_path, **({"gtin": gtin(row[10])} if gtin(row[10]) else {})})
            sizes.append(size)
        chosen = min([v for v in variants if v["available"]] or variants, key=lambda v: v["price"])
        is_mtb = any(token in f"{category} {name}".casefold() for token in ("mountain", "mtb", "siskiu", "xtrada", "syncline", "premier", "heist"))
        products.append(product(
            product_id=930000 + offset, name=f"Polygon {name} {year} - {colour}", brand="Polygon", model=clean(rows[0][10]) or name,
            price=chosen["price"], mrp=chosen["mrp"], stock=sum(v["stockQuantity"] for v in variants), images=[image_path],
            page="bikes-frames.html", label="Bikes & Frames", main="mountain-bikes" if is_mtb else "road-bikes", sub=category.casefold(),
            description=f"Polygon {name} {year} in {colour}. Choose an available size before checkout.",
            specs=["Brand : Polygon", f"Year : {year}", f"Category : {category}", f"Colour : {colour}"],
            variants=variants, sizes=list(dict.fromkeys(sizes)), colours=[colour], source_url=source_url))


def merida_official_image(model, year, cache, audit):
    query_model = re.sub(r"\s*-?\s*(?:KIT-)?FRAME SET.*$", "", clean(model), flags=re.I).strip()
    cache_key = norm(query_model)
    disk_key = f"Merida|{cache_key}"
    if disk_key in SOURCE_CACHE and (ROOT / SOURCE_CACHE[disk_key]["image"]).exists():
        saved = SOURCE_CACHE[disk_key]
        audit["sourcedImages"].append({"brand": "Merida", "model": model, "sourcePage": saved["sourcePage"], "sourceName": "Merida official"})
        return saved["image"], saved["sourcePage"]
    if cache_key in cache:
        return cache[cache_key]
    alternate = re.sub(r"\bRIM\b", "", query_model, flags=re.I)
    family = " ".join(query_model.split()[:2]) if query_model.casefold().startswith("scultura endurance") else query_model.split()[0]
    queries = list(dict.fromkeys([query_model, clean(alternate), family]))
    exact = []
    for query in queries:
        url = "https://www.merida-bikes.com/en/bikefinder?" + urllib.parse.urlencode({"query": query})
        page = fetch(url).decode("utf-8", "ignore")
        cards = re.findall(r'<div class="bike-card".*?<img[^>]+src="([^"]+)"[^>]+alt="([^"]+)".*?<a href="([^"]+)"[^>]*data-hover', page, re.S | re.I)
        desired = {cache_key, norm(alternate)}
        exact = [card for card in cards if norm(html.unescape(card[1])) in desired]
        if not exact and cards:
            model_tokens = set(norm(query_model).split())
            exact = sorted(cards, key=lambda card: -len(model_tokens & set(norm(html.unescape(card[1])).split())))
        if exact:
            break
    if not exact:
        raise RuntimeError(f"No official Merida product match for {model}")
    year_text = clean(year)
    chosen = next((card for card in exact if f"/{year_text}/" in card[0]), exact[0])
    image_url = html.unescape(chosen[0])
    source_url = urllib.parse.urljoin("https://www.merida-bikes.com", chosen[2])
    result = (save_remote(image_url), source_url)
    cache[cache_key] = result
    SOURCE_CACHE[disk_key] = {"image": result[0], "sourcePage": result[1]}
    CACHE_PATH.write_text(json.dumps(SOURCE_CACHE, indent=2) + "\n", encoding="utf-8")
    audit["sourcedImages"].append({"brand": "Merida", "model": model, "sourcePage": source_url, "sourceName": "Merida official"})
    return result


def import_merida(path, products, audit):
    sheet = openpyxl.load_workbook(path, data_only=True)["Stock Sheet"]
    grouped = defaultdict(list)
    for row in range(3, 278):
        model = clean(sheet.cell(row, 5).value)
        if not model:
            continue
        group_key = (clean(sheet.cell(row, 4).value), norm(model), norm(sheet.cell(row, 6).value))
        grouped[group_key].append(row)
    image_cache = {}
    for offset, rows in enumerate(grouped.values(), 1):
        first = rows[0]
        year, model, colour = clean(sheet.cell(first, 4).value), clean(sheet.cell(first, 5).value), clean(sheet.cell(first, 6).value) or "Standard"
        print(f"Merida {offset}/{len(grouped)}: {model}", flush=True)
        image_path, source_url = merida_official_image(model, year, image_cache, audit)
        variants, sizes = [], []
        for row in rows:
            size = clean(sheet.cell(row, 7).value) or "Standard"
            qty = quantity(sheet.cell(row, 8).value)
            mrp = money(sheet.cell(row, 10).value)
            price = money(sheet.cell(row, 11).value) or mrp
            variants.append({"key": key_for("merida", year, model, colour, size), "size": size, "colour": colour, "build": "", "price": price, "mrp": max(mrp, price), "available": qty > 0, "stockStatus": "available" if qty > 0 else "coming-soon", "stockQuantity": qty, "image": image_path, **({"gtin": gtin(sheet.cell(row, 3).value)} if gtin(sheet.cell(row, 3).value) else {})})
            sizes.append(size)
        chosen = min([v for v in variants if v["available"]] or variants, key=lambda v: v["price"])
        lowered = model.casefold()
        is_frame = "frame" in lowered
        is_mtb = any(token in lowered for token in ("big nine", "big seven", "one-twenty"))
        products.append(product(
            product_id=931000 + offset, name=f"Merida {model} {year} - {colour}", brand="Merida", model=clean(sheet.cell(first, 2).value) or model,
            price=chosen["price"], mrp=chosen["mrp"], stock=sum(v["stockQuantity"] for v in variants), images=[image_path],
            page="bikes-frames.html", label="Bikes & Frames", main="frames" if is_frame else ("mountain-bikes" if is_mtb else "road-and-hybrid-bikes"),
            sub="frames" if is_frame else "complete bicycles",
            description=f"Merida {model} {year} in {colour}. Choose an available size before checkout.",
            specs=["Brand : Merida", f"Year : {year}", f"Colour : {colour}"],
            variants=variants, sizes=list(dict.fromkeys(sizes)), colours=[colour], source_url=source_url))


def main():
    missing = [str(path) for path in SOURCES.values() if not path.exists()]
    if missing:
        raise SystemExit("Missing workbook(s): " + ", ".join(missing))
    products = []
    audit = {"held": [], "sourcedImages": [], "workbookImages": {}}
    import_tifosi(SOURCES["tifosi"], products, audit)
    import_ravemen(SOURCES["ravemen"], products, audit)
    import_safety(SOURCES["safety"], products, audit)
    import_tacx(SOURCES["tacx"], products, audit)
    import_polygon(SOURCES["polygon"], products, audit)
    import_merida(SOURCES["merida"], products, audit)
    DATA.parent.mkdir(parents=True, exist_ok=True)
    DATA.write_text(json.dumps({"version": 1, "stockAsOf": "2026-10-16", "products": products, "audit": audit}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    allowlist_path = ROOT / "assets" / "data" / "product-allowlist.json"
    allowlist = json.loads(allowlist_path.read_text(encoding="utf-8"))
    allowlist["productIds"] = sorted(set(map(str, allowlist["productIds"])) | {str(p["id"]) for p in products}, key=int)
    allowlist_path.write_text(json.dumps(allowlist, indent=2) + "\n", encoding="utf-8")
    brands = sorted({p["brand"] for p in products})
    print(json.dumps({"products": len(products), "available": sum(p["available"] for p in products), "comingSoon": sum(not p["available"] for p in products), "byBrand": {brand: sum(p["brand"] == brand for p in products) for brand in brands}, "sourcedImages": len(audit["sourcedImages"]), "held": len(audit["held"])}, indent=2))


if __name__ == "__main__":
    main()
